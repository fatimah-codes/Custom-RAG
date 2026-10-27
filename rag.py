"""
DocuMind AI  -  Multi-format RAG Agent + Custom ChatGPT
=======================================================
Run with:   streamlit run rag.py

Features
- API-key gate first, console only after the key is verified
- Upload PDF, Word (.docx), Excel (.xlsx/.xls), CSV, TXT/MD and images (PNG/JPG/WEBP)
  (images are read with OpenAI vision: OCR + description)
- Multi-file knowledge base with FAISS vector search (similarity or MMR)
- Smart answers: from your documents when possible; otherwise the bot says
  "This isn't in the document you provided, but the answer is: ..." and answers anyway
- Free-chat mode (plain ChatGPT), conversation memory, follow-up question rewriting
- Streaming answers, source citations, quick actions (summary, key points, ...)
- Chat history in the sidebar: new chat, switch between past chats, delete (kept for the browser session)
- 6 switchable themes + animation toggle, chat export, library & stats tab
"""
from __future__ import annotations

import base64
import io
import os
import re
import tempfile
import time
from datetime import datetime

import streamlit as st
from langchain_community.document_loaders import PyPDFLoader
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.messages import AIMessage, HumanMessage
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

# ---------------------------------------------------
# Page setup
# ---------------------------------------------------
st.set_page_config(page_title="DocuMind AI", page_icon="🧠", layout="wide")

ss = st.session_state
ss.setdefault("theme", "Aurora Night")
ss.setdefault("raw_cache", {})
ss.setdefault("chats", {})   # {chat_id: {"title": str, "messages": list, "created": datetime}}


def new_chat() -> None:
    """Start a fresh conversation (re-uses the current one if it is still empty)."""
    cur = ss.chats.get(ss.get("active"))
    if cur is not None and not cur["messages"]:
        ss.messages = cur["messages"]
        return
    cid = datetime.now().strftime("%Y%m%d%H%M%S%f")
    ss.chats[cid] = {"title": "New chat", "messages": [], "created": datetime.now()}
    ss.active = cid
    ss.messages = ss.chats[cid]["messages"]


def switch_chat(cid: str) -> None:
    if cid in ss.chats:
        ss.active = cid
        ss.messages = ss.chats[cid]["messages"]


def delete_chat(cid: str) -> None:
    ss.chats.pop(cid, None)
    if ss.get("active") == cid or ss.get("active") not in ss.chats:
        ss.active = None
        if ss.chats:
            switch_chat(list(ss.chats)[-1])
        else:
            new_chat()


if ss.get("active") not in ss.chats:
    new_chat()
ss.messages = ss.chats[ss.active]["messages"]

MAX_TABLE_ROWS = 3000
IMAGE_EXT = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp"}
ALLOWED_TYPES = ["pdf", "docx", "xlsx", "xls", "csv", "txt", "md", "png", "jpg", "jpeg", "webp"]

# ---------------------------------------------------
# Themes
# ---------------------------------------------------
THEMES = {
    "Aurora Night": dict(
        bg=["#1b0a3c", "#3a1c71", "#0f4c81", "#0b7a75"], text="#f4f1ff", muted="rgba(244,241,255,.7)",
        card="rgba(255,255,255,.08)", border="rgba(255,255,255,.18)", a1="#ff6ec7", a2="#7c4dff", a3="#4facfe",
        input="rgba(255,255,255,.12)", sidebar="rgba(15,8,40,.70)", btn="#ffffff"),
    "Sunset Blaze": dict(
        bg=["#2b0f1f", "#6b1d3a", "#b4361b", "#d98a12"], text="#fff7ed", muted="rgba(255,247,237,.72)",
        card="rgba(255,255,255,.09)", border="rgba(255,255,255,.20)", a1="#ff512f", a2="#f09819", a3="#ff5e98",
        input="rgba(255,255,255,.13)", sidebar="rgba(43,15,31,.72)", btn="#ffffff"),
    "Ocean Breeze": dict(
        bg=["#031d44", "#04395e", "#0a6c74", "#0e8f8f"], text="#e8fbff", muted="rgba(232,251,255,.72)",
        card="rgba(255,255,255,.08)", border="rgba(255,255,255,.18)", a1="#00d2ff", a2="#3a7bd5", a3="#00f2c3",
        input="rgba(255,255,255,.12)", sidebar="rgba(3,29,68,.72)", btn="#ffffff"),
    "Emerald Forest": dict(
        bg=["#04200f", "#0b4f2c", "#136f4f", "#1d8a6e"], text="#effff4", muted="rgba(239,255,244,.72)",
        card="rgba(255,255,255,.08)", border="rgba(255,255,255,.18)", a1="#43e97b", a2="#1db954", a3="#38f9d7",
        input="rgba(255,255,255,.12)", sidebar="rgba(4,32,15,.72)", btn="#04200f"),
    "Cyber Neon": dict(
        bg=["#07040f", "#12062e", "#2a0a4a", "#001a33"], text="#e8fbff", muted="rgba(232,251,255,.7)",
        card="rgba(255,255,255,.06)", border="rgba(0,255,240,.30)", a1="#00fff0", a2="#ff00e5", a3="#8a2be2",
        input="rgba(255,255,255,.08)", sidebar="rgba(7,4,15,.80)", btn="#07040f"),
    "Daylight": dict(
        bg=["#fdf2ff", "#e0f2fe", "#fef9c3", "#dcfce7"], text="#1f1b3a", muted="rgba(31,27,58,.68)",
        card="rgba(255,255,255,.72)", border="rgba(80,60,160,.20)", a1="#7c4dff", a2="#ec4899", a3="#0ea5e9",
        input="rgba(255,255,255,.95)", sidebar="rgba(255,255,255,.78)", btn="#ffffff"),
}

CSS_TEMPLATE = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Poppins:wght@300;400;500;600;700&display=swap');
:root{
  --text:@@TEXT@@; --muted:@@MUTED@@; --card:@@CARD@@; --border:@@BORDER@@;
  --a1:@@A1@@; --a2:@@A2@@; --a3:@@A3@@; --input:@@INPUT@@; --btn:@@BTN@@;
}
html, body, [class*="css"], .stApp { font-family:'Poppins',sans-serif; }

/* animated gradient background */
.stApp{
  background:linear-gradient(-45deg,@@BG1@@,@@BG2@@,@@BG3@@,@@BG4@@);
  background-size:400% 400%; animation:gradientShift 22s ease infinite; color:var(--text);
}
@keyframes gradientShift{0%{background-position:0% 50%}50%{background-position:100% 50%}100%{background-position:0% 50%}}

/* floating orbs */
.orb{position:fixed;border-radius:50%;filter:blur(70px);opacity:.40;z-index:0;pointer-events:none;
     animation:floaty 15s ease-in-out infinite}
.orb.o1{width:340px;height:340px;background:var(--a1);top:-90px;left:-70px}
.orb.o2{width:400px;height:400px;background:var(--a3);bottom:-130px;right:-90px;animation-delay:-4s}
.orb.o3{width:260px;height:260px;background:var(--a2);top:45%;left:55%;animation-delay:-8s;opacity:.28}
@keyframes floaty{0%,100%{transform:translate(0,0) scale(1)}33%{transform:translate(40px,-50px) scale(1.12)}66%{transform:translate(-30px,40px) scale(.94)}}

header[data-testid="stHeader"]{background:transparent}
#MainMenu, footer{visibility:hidden}
.block-container{position:relative;z-index:1;padding-top:2rem;animation:fadeUp .8s ease both}
@keyframes fadeUp{from{opacity:0;transform:translateY(24px)}to{opacity:1;transform:translateY(0)}}

/* hero */
.hero{text-align:center;margin:.5rem 0 1.6rem;animation:fadeUp 1s ease both}
.hero h1{font-size:3rem;font-weight:700;margin:0;
  background:linear-gradient(90deg,var(--a1),var(--a2),var(--a3),var(--a1));background-size:300% 100%;
  -webkit-background-clip:text;background-clip:text;-webkit-text-fill-color:transparent;animation:shine 6s linear infinite}
@keyframes shine{to{background-position:300% 0}}
.hero p{color:var(--muted);font-weight:300;letter-spacing:.5px;margin-top:.3rem}
.emoji{font-size:3rem;text-align:center;line-height:1.2}
.bounce{display:inline-block;animation:bounce 2.4s ease-in-out infinite}
@keyframes bounce{0%,100%{transform:translateY(0)}50%{transform:translateY(-10px)}}

/* glass cards */
[data-testid="stForm"],[data-testid="stFileUploader"],[data-testid="stExpander"],[data-testid="stMetric"]{
  background:var(--card);backdrop-filter:blur(14px);border:1px solid var(--border);border-radius:20px;
  box-shadow:0 10px 40px rgba(0,0,0,.22);padding:.6rem;animation:fadeUp .9s ease both}
[data-testid="stForm"]:hover,[data-testid="stFileUploader"]:hover,[data-testid="stMetric"]:hover{
  box-shadow:0 14px 50px color-mix(in srgb,var(--a2) 45%,transparent);transition:box-shadow .4s ease}

.feat-row{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:1.2rem 0}
.feat{background:var(--card);border:1px solid var(--border);border-radius:18px;padding:.9rem;text-align:center;
      backdrop-filter:blur(10px);transition:transform .3s ease, box-shadow .3s ease;animation:fadeUp 1.2s ease both}
.feat:hover{transform:translateY(-6px) rotate(-1deg);box-shadow:0 14px 30px color-mix(in srgb,var(--a2) 40%,transparent)}
.feat b{display:block;margin-top:.2rem;font-size:.9rem}
.feat span{font-size:1.8rem}

/* sidebar */
[data-testid="stSidebar"]{background:@@SIDEBAR@@;backdrop-filter:blur(18px);border-right:1px solid var(--border)}

/* text colours (targeted, so dropdown popovers keep their defaults) */
[data-testid="stWidgetLabel"],[data-testid="stWidgetLabel"] p,[data-testid="stMarkdownContainer"] p,
[data-testid="stMarkdownContainer"] li,[data-testid="stMarkdownContainer"] h1,[data-testid="stMarkdownContainer"] h2,
[data-testid="stMarkdownContainer"] h3,[data-testid="stMarkdownContainer"] h4,[data-testid="stCaptionContainer"],
[data-testid="stMetricValue"],[data-testid="stMetricLabel"],[data-testid="stMetricLabel"] p,
[data-testid="stRadio"] label,[data-testid="stCheckbox"] label{color:var(--text)!important}
[data-testid="stCaptionContainer"]{color:var(--muted)!important}
[data-testid="stThumbValue"],[data-testid="stTickBarMin"],[data-testid="stTickBarMax"]{color:var(--text)!important;font-weight:600}

/* inputs */
.stTextInput input{background:var(--input)!important;color:var(--text)!important;border:1px solid var(--border)!important;
  border-radius:12px!important;transition:all .3s ease}
.stTextInput input:focus{border-color:var(--a2)!important;box-shadow:0 0 0 3px color-mix(in srgb,var(--a2) 35%,transparent)!important}
[data-baseweb="select"]>div{background:var(--input)!important;border:1px solid var(--border)!important;border-radius:12px!important}
[data-baseweb="select"] *{color:var(--text)!important}
[data-baseweb="select"] svg{fill:var(--text)!important}

/* buttons */
.stButton>button,[data-testid="stFormSubmitButton"]>button,[data-testid="stDownloadButton"]>button{
  background:linear-gradient(135deg,var(--a1),var(--a2),var(--a3));background-size:200% 200%;border:none;border-radius:14px;
  padding:.55rem 1.2rem;width:100%;transition:transform .25s ease,box-shadow .25s ease,background-position .6s ease}
.stButton>button p,[data-testid="stFormSubmitButton"]>button p,[data-testid="stDownloadButton"]>button p{
  color:var(--btn)!important;font-weight:600}
.stButton>button:hover,[data-testid="stFormSubmitButton"]>button:hover,[data-testid="stDownloadButton"]>button:hover{
  transform:translateY(-3px) scale(1.02);box-shadow:0 10px 28px color-mix(in srgb,var(--a2) 55%,transparent);background-position:100% 0}
.stButton>button:disabled{opacity:.45;transform:none;box-shadow:none}

/* uploader */
[data-testid="stFileUploaderDropzone"]{background:var(--input)!important;border:2px dashed var(--a2)!important;border-radius:16px!important}
[data-testid="stFileUploaderDropzone"] *,[data-testid="stFileUploader"] label,[data-testid="stFileUploader"] small{color:var(--text)!important}
[data-testid="stFileUploaderDropzone"] button{background:linear-gradient(135deg,var(--a1),var(--a2))!important;border:none!important;border-radius:10px!important}
[data-testid="stFileUploaderDropzone"] button *{color:var(--btn)!important}

/* tabs */
button[data-baseweb="tab"] p{color:var(--text)!important;font-weight:500}
[data-baseweb="tab-highlight"]{background:linear-gradient(90deg,var(--a1),var(--a3))!important;height:3px!important}

/* chat */
[data-testid="stChatMessage"]{background:var(--card);border:1px solid var(--border);border-radius:18px;padding:1rem;
  margin-bottom:.8rem;backdrop-filter:blur(10px);animation:slideIn .5s cubic-bezier(.2,.8,.2,1) both}
[data-testid="stChatMessage"] *{color:var(--text)}
@keyframes slideIn{from{opacity:0;transform:translateX(-30px) scale(.97)}to{opacity:1;transform:translateX(0) scale(1)}}
[data-testid="stBottom"],[data-testid="stBottom"]>div{background:transparent!important}
[data-testid="stChatInput"]{border-radius:16px;background:var(--card)!important;border:1px solid var(--border)}
[data-testid="stChatInput"] textarea{color:var(--text)!important;background:transparent!important}

/* alerts */
[data-testid="stAlert"]{background:var(--card)!important;border:1px solid var(--border)!important;border-radius:14px;animation:fadeUp .6s ease both}
[data-testid="stAlert"] *{color:var(--text)!important}

/* badges */
.badge,.tag{display:inline-block;padding:.25rem .8rem;border-radius:999px;font-size:.78rem;font-weight:500;color:var(--text)}
.badge{background:rgba(67,233,123,.18);border:1px solid rgba(67,233,123,.6);animation:pulse 2s ease-in-out infinite}
.tag.doc{background:rgba(67,233,123,.18);border:1px solid rgba(67,233,123,.6)}
.tag.gen{background:rgba(255,193,7,.18);border:1px solid rgba(255,193,7,.7)}
.tag.chat{background:rgba(79,172,254,.18);border:1px solid rgba(79,172,254,.7)}
@keyframes pulse{0%,100%{box-shadow:0 0 0 0 rgba(67,233,123,.5)}50%{box-shadow:0 0 0 10px rgba(67,233,123,0)}}

.src-title{font-weight:600;color:var(--a3)}
.filechip{display:inline-block;margin:.15rem .25rem .15rem 0;padding:.25rem .75rem;border-radius:999px;
  background:var(--card);border:1px solid var(--border);font-size:.82rem}
</style>
<div class="orb o1"></div><div class="orb o2"></div><div class="orb o3"></div>
"""


def inject_theme(name: str, animate: bool = True) -> None:
    t = THEMES[name]
    css = CSS_TEMPLATE
    repl = {
        "TEXT": t["text"], "MUTED": t["muted"], "CARD": t["card"], "BORDER": t["border"],
        "A1": t["a1"], "A2": t["a2"], "A3": t["a3"], "INPUT": t["input"], "BTN": t["btn"],
        "SIDEBAR": t["sidebar"], "BG1": t["bg"][0], "BG2": t["bg"][1], "BG3": t["bg"][2], "BG4": t["bg"][3],
    }
    for k, v in repl.items():
        css = css.replace(f"@@{k}@@", v)
    if not animate:
        css = css.replace("</style>", "*{animation:none!important;transition:none!important}</style>", 1)
    st.markdown(css, unsafe_allow_html=True)


def _set_theme(widget_key: str) -> None:
    ss.theme = ss[widget_key]


def theme_picker(where: str) -> None:
    key = f"theme_widget_{where}"
    names = list(THEMES)
    st.selectbox("🎨 Theme", names, index=names.index(ss.theme), key=key,
                 on_change=_set_theme, args=(key,))


inject_theme(ss.theme, ss.get("anim_w", True))


# ---------------------------------------------------
# Step 1: API key gate
# ---------------------------------------------------
def validate_key(key: str):
    from openai import OpenAI
    try:
        OpenAI(api_key=key).models.list()
        return True, ""
    except Exception as e:  # noqa: BLE001
        msg = str(e)
        if "401" in msg or "invalid_api_key" in msg or "Incorrect API key" in msg:
            return False, "This API key is invalid. Please check it and try again."
        return False, f"Could not verify the key: {msg[:200]}"


def api_key_screen() -> None:
    st.markdown("<style>[data-testid='stSidebar']{display:none;}</style>", unsafe_allow_html=True)
    st.markdown(
        """
        <div class="hero">
            <div class="emoji bounce">🧠</div>
            <h1>DocuMind AI</h1>
            <p>Chat with your PDFs, Word, Excel &amp; images — with a ChatGPT brain for everything else</p>
        </div>
        <div class="feat-row">
            <div class="feat"><span>📄</span><b>PDF &amp; Word</b></div>
            <div class="feat"><span>📊</span><b>Excel &amp; CSV</b></div>
            <div class="feat"><span>🖼️</span><b>Image reading</b></div>
            <div class="feat"><span>💬</span><b>Smart ChatGPT fallback</b></div>
            <div class="feat"><span>🎨</span><b>6 themes</b></div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    _, mid, _ = st.columns([1, 2, 1])
    with mid:
        theme_picker("gate")
        with st.form("api_key_form"):
            key = st.text_input("🔐 OpenAI API Key", type="password", placeholder="sk-...")
            go = st.form_submit_button("🚀 Unlock the console")
        st.caption("Your key stays in this browser session only and is never saved to disk.")
        if go:
            if not key.strip():
                st.warning("Please enter your API key first.")
            else:
                with st.spinner("Verifying your key…"):
                    ok, msg = validate_key(key.strip())
                if ok:
                    ss.api_key = key.strip()
                    ss.key_ok = True
                    st.balloons()
                    time.sleep(1.2)
                    st.rerun()
                else:
                    st.error(msg)


if not ss.get("key_ok"):
    api_key_screen()
    st.stop()

os.environ["OPENAI_API_KEY"] = ss.api_key


# ---------------------------------------------------
# File loaders (PDF, DOCX, Excel/CSV, TXT, images)
# ---------------------------------------------------
def load_pdf(name: str, data: bytes):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(data)
        path = tmp.name
    try:
        pages = PyPDFLoader(path).load()
    finally:
        os.unlink(path)
    docs = []
    for d in pages:
        if d.page_content.strip():
            pg = d.metadata.get("page")
            d.metadata = {"source": name, "kind": "pdf", "loc": f"p. {pg + 1}" if isinstance(pg, int) else "p. ?"}
            docs.append(d)
    warns = [] if docs else [f"{name}: no selectable text found (scanned PDF?). Upload its pages as images instead."]
    return docs, warns


def load_docx(name: str, data: bytes):
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    docs = []
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if text.strip():
        docs.append(Document(page_content=text, metadata={"source": name, "kind": "docx", "loc": "body"}))
    for i, table in enumerate(doc.tables, 1):
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        if any(rows):
            docs.append(Document(page_content="\n".join(rows),
                                 metadata={"source": name, "kind": "docx", "loc": f"table {i}"}))
    return docs, ([] if docs else [f"{name}: the document looks empty."])


def load_table(name: str, data: bytes, ext: str):
    import pandas as pd

    if ext == "csv":
        sheets = {"csv": pd.read_csv(io.BytesIO(data), encoding_errors="replace")}
    else:
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    docs, warns = [], []
    for sheet, df in sheets.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        if df.empty:
            continue
        if len(df) > MAX_TABLE_ROWS:
            warns.append(f"{name} [{sheet}]: only the first {MAX_TABLE_ROWS} rows were indexed.")
            df = df.head(MAX_TABLE_ROWS)
        cols = list(df.columns)
        lines = [f"Sheet: {sheet}", "Columns: " + ", ".join(map(str, cols))]
        for i, row in enumerate(df.itertuples(index=False), start=2):
            cells = "; ".join(f"{c}={v}" for c, v in zip(cols, row) if pd.notna(v))
            lines.append(f"Row {i}: {cells}")
        docs.append(Document(page_content="\n".join(lines),
                             metadata={"source": name, "kind": "table", "loc": f"sheet {sheet}"}))
    return docs, warns


def load_text(name: str, data: bytes):
    text = data.decode("utf-8", errors="replace")
    docs = [Document(page_content=text, metadata={"source": name, "kind": "text", "loc": "text"})] if text.strip() else []
    return docs, ([] if docs else [f"{name}: the file is empty."])


def load_image(name: str, data: bytes, ext: str, vision_model: str):
    b64 = base64.b64encode(data).decode()
    msg = HumanMessage(content=[
        {"type": "text", "text": (
            "You are an OCR + image analysis engine. First transcribe ALL visible text exactly "
            "(keep tables as rows separated by ' | '). Then add a section 'Description:' with a "
            "detailed description of the image: objects, people, charts, numbers, colours, layout.")},
        {"type": "image_url", "image_url": {"url": f"data:{IMAGE_EXT[ext]};base64,{b64}"}},
    ])
    out = ChatOpenAI(model=vision_model, temperature=0).invoke([msg]).content
    return [Document(page_content=out, metadata={"source": name, "kind": "image", "loc": "image"})], []


def load_file(f, vision_model: str):
    name, data = f.name, f.getvalue()
    ext = name.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return load_pdf(name, data)
    if ext == "docx":
        return load_docx(name, data)
    if ext in ("xlsx", "xls", "csv"):
        return load_table(name, data, ext)
    if ext in ("txt", "md"):
        return load_text(name, data)
    if ext in IMAGE_EXT:
        return load_image(name, data, ext, vision_model)
    return [], [f"{name}: unsupported file type."]


def build_index(files, size: int, overlap: int, vision_model: str):
    all_docs, warnings = [], []
    with st.status("Reading and indexing your files…", expanded=True) as status:
        for f in files:
            key = (f.name, f.size)
            if key not in ss.raw_cache:
                st.write(f"📥 Reading **{f.name}**")
                try:
                    ss.raw_cache[key] = load_file(f, vision_model)
                except Exception as e:  # noqa: BLE001
                    ss.raw_cache[key] = ([], [f"{f.name}: could not be read ({str(e)[:150]})"])
            docs, warns = ss.raw_cache[key]
            all_docs += docs
            warnings += warns
        if not all_docs:
            status.update(label="Nothing readable found", state="error")
            return None, [], warnings
        st.write("✂️ Splitting into chunks")
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=size, chunk_overlap=overlap, separators=["\n\n", "\n", ". ", " ", ""])
        chunks = splitter.split_documents(all_docs)
        st.write(f"🧬 Creating embeddings for {len(chunks)} chunks")
        vs = FAISS.from_documents(chunks, OpenAIEmbeddings(model="text-embedding-3-small"))
        status.update(label=f"Indexed {len(files)} file(s) · {len(chunks)} chunks", state="complete", expanded=False)
    return vs, chunks, warnings


# ---------------------------------------------------
# Sidebar
# ---------------------------------------------------
with st.sidebar:
    st.markdown("### 🕘 Chat history")
    st.button("➕ New chat", key="new_chat_btn", on_click=new_chat)
    hist_box = st.container()   # filled at the end of the script, once titles are up to date
    st.divider()
    st.markdown("### ⚙️ Control Center")
    st.markdown('<span class="badge">● API key connected</span>', unsafe_allow_html=True)
    st.write("")
    theme_picker("sidebar")
    st.checkbox("✨ Animations", value=True, key="anim_w")
    mode = st.radio("🧭 Answer mode",
                    ["Smart (documents + ChatGPT fallback)", "Free chat (ChatGPT only)"])
    model_name = st.selectbox("Chat model", ["gpt-4o-mini", "gpt-4o", "gpt-4.1-mini", "gpt-4.1"])
    temperature = st.slider("Creativity (temperature)", 0.0, 1.0, 0.2, step=0.1)
    with st.expander("🔎 Retrieval settings"):
        chunk_size = st.slider("Chunk size", 500, 2000, 1000, step=100)
        chunk_overlap = st.slider("Chunk overlap", 0, 400, 150, step=50)
        top_k = st.slider("Retrieved chunks (k)", 2, 10, 4)
        use_mmr = st.checkbox("Diverse results (MMR)", value=False)
    st.divider()
    if st.button("🧹 Clear current chat"):
        ss.messages.clear()
        ss.chats[ss.active]["title"] = "New chat"
        st.rerun()
    if ss.messages:
        export = ["# DocuMind AI – chat export", f"_{datetime.now():%Y-%m-%d %H:%M}_", ""]
        for m in ss.messages:
            who = "You" if m["role"] == "user" else "DocuMind"
            export += [f"**{who}:** {m['content']}", ""]
        st.download_button("⬇️ Export chat (.md)", "\n".join(export), file_name="documind_chat.md",
                           mime="text/markdown")
    if st.button("🔑 Change API key"):
        theme = ss.theme
        ss.clear()
        ss.theme = theme
        st.rerun()

# ---------------------------------------------------
# Hero + uploader + indexing
# ---------------------------------------------------
st.markdown(
    """
    <div class="hero">
        <div class="emoji bounce">🧠</div>
        <h1>DocuMind AI</h1>
        <p>Upload PDFs, Word, Excel, CSV, text or images — ask anything. Beyond your files, I still answer like ChatGPT.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

uploaded = st.file_uploader(
    "📎 Drop your files here (PDF · DOCX · XLSX/XLS · CSV · TXT/MD · PNG/JPG/WEBP)",
    type=ALLOWED_TYPES, accept_multiple_files=True)
st.caption("Tip: old .doc files are not supported — save them as .docx first.")

if uploaded:
    sig = tuple((f.name, f.size) for f in uploaded) + (chunk_size, chunk_overlap)
    if ss.get("file_sig") != sig:
        vs, chunks, warns = build_index(uploaded, chunk_size, chunk_overlap, model_name)
        ss.vectorstore, ss.chunks, ss.warns, ss.file_sig = vs, chunks, warns, sig
        if ss.get("file_names") is not None and ss.file_names != [f.name for f in uploaded] and ss.messages:
            new_chat()   # different files -> fresh conversation, old one stays in history
        ss.file_names = [f.name for f in uploaded]
    for w in ss.get("warns", []):
        st.warning(w)
else:
    for k in ("vectorstore", "chunks", "file_sig", "file_names", "warns"):
        ss.pop(k, None)

has_docs = ss.get("vectorstore") is not None

# ---------------------------------------------------
# RAG prompt + helpers
# ---------------------------------------------------
SYSTEM_PROMPT = (
    "You are DocuMind, a friendly, sharp AI assistant (a custom ChatGPT) that works with the user's "
    "uploaded documents.\n\n"
    "CONTEXT FROM THE USER'S DOCUMENTS:\n{context}\n\n"
    "STRICT OUTPUT FORMAT - begin EVERY reply with exactly one tag, then the answer:\n"
    "[DOC]     The context contains the answer. Answer using ONLY the context and cite sources in "
    "parentheses, e.g. (report.pdf, p. 3).\n"
    "[GENERAL] The user asked a real question but the context does NOT contain the answer. Right after the tag "
    "write exactly: \"⚠️ This isn't in the document you provided, but the answer is: \" and then give a correct, "
    "helpful answer from your general knowledge.\n"
    "[CHAT]    Greetings, thanks, small talk, questions about you, or whenever the context says "
    "NO DOCUMENTS or FREE CHAT. Reply naturally with no warning line.\n\n"
    "Never invent facts about the documents. Use markdown formatting. Be concise but complete. "
    "Reply in the same language as the user."
)
PROMPT = ChatPromptTemplate.from_messages([
    ("system", SYSTEM_PROMPT),
    MessagesPlaceholder("history"),
    ("human", "{question}"),
])

QUICK_ACTIONS = {
    "📝 Summarize": "Give a clear, well-structured summary of the uploaded document(s).",
    "🔑 Key points": "List the key points, facts and takeaways from the uploaded document(s) as bullet points.",
    "❓ Suggest questions": "Suggest 6 interesting questions I could ask about the uploaded document(s).",
    "🧒 Explain simply": "Explain the main idea of the uploaded document(s) in very simple terms for a beginner.",
}

MODE_BADGES = {
    "DOC": '<span class="tag doc">📄 From your documents</span>',
    "GENERAL": '<span class="tag gen">🌐 Not in your documents · general knowledge</span>',
    "CHAT": '<span class="tag chat">💬 Chat</span>',
}


def doc_label(d: Document) -> str:
    return f"{d.metadata.get('source', '?')} · {d.metadata.get('loc', '')}".strip(" ·")


def format_context(docs) -> str:
    return "\n\n".join(f"[Source: {doc_label(d)}]\n{d.page_content}" for d in docs)


def sample_chunks(chunks, n: int = 14):
    if len(chunks) <= n:
        return chunks
    step = len(chunks) / n
    return [chunks[int(i * step)] for i in range(n)]


def retrieve(query: str, k: int, mmr: bool):
    kwargs = {"k": k, "fetch_k": max(20, k * 4)} if mmr else {"k": k}
    retriever = ss.vectorstore.as_retriever(search_type="mmr" if mmr else "similarity", search_kwargs=kwargs)
    return retriever.invoke(query)


def standalone_question(llm, question: str, history) -> str:
    if not history:
        return question
    convo = "\n".join(f"{m['role']}: {m['content'][:600]}" for m in history[-6:])
    p = ChatPromptTemplate.from_template(
        "Rewrite the follow-up as a standalone search query using the conversation for context. "
        "Return ONLY the rewritten question.\n\nConversation:\n{convo}\n\nFollow-up: {q}\n\nStandalone:")
    try:
        return (p | llm | StrOutputParser()).invoke({"convo": convo, "q": question}).strip() or question
    except Exception:  # noqa: BLE001
        return question


def stream_answer(chain, inputs, on_mode):
    """Yield the answer text, strip the leading [DOC]/[GENERAL]/[CHAT] tag and report it via on_mode."""
    buf, decided = "", False
    for chunk in chain.stream(inputs):
        if decided:
            yield chunk
            continue
        buf += chunk
        m = re.match(r"\s*\[(DOC|GENERAL|CHAT)\]\s*", buf)
        if m:
            on_mode(m.group(1))
            decided = True
            rest = buf[m.end():]
            if rest:
                yield rest
        elif len(buf) > 24:
            on_mode("DOC" if has_docs else "CHAT")
            decided = True
            yield buf
    if not decided and buf:
        on_mode("DOC" if has_docs else "CHAT")
        yield buf


def render_sources(sources) -> None:
    if not sources:
        return
    with st.expander(f"🔍 Sources ({len(sources)} chunks)"):
        for s in sources:
            st.markdown(f'<div class="src-title">{s["label"]}</div>', unsafe_allow_html=True)
            st.text(s["text"])
            st.divider()


def render_message(m) -> None:
    with st.chat_message(m["role"], avatar="🧑‍💻" if m["role"] == "user" else "🧠"):
        if m["role"] == "assistant" and m.get("mode"):
            st.markdown(MODE_BADGES[m["mode"]], unsafe_allow_html=True)
        st.markdown(m["content"])
        render_sources(m.get("sources"))


def handle_question(question: str, wide: bool = False) -> None:
    history = [m for m in ss.messages if m["role"] in ("user", "assistant")][-6:]
    lc_history = [HumanMessage(m["content"]) if m["role"] == "user" else AIMessage(m["content"]) for m in history]

    ss.messages.append({"role": "user", "content": question})
    if ss.chats[ss.active]["title"] == "New chat":
        ss.chats[ss.active]["title"] = question.strip().replace("\n", " ")[:40]
    render_message(ss.messages[-1])

    llm = ChatOpenAI(model=model_name, temperature=temperature, streaming=True)
    use_docs = has_docs and mode.startswith("Smart")

    with st.chat_message("assistant", avatar="🧠"):
        badge_ph = st.empty()
        try:
            docs = []
            if use_docs:
                with st.spinner("Searching your documents…"):
                    if wide:
                        docs = sample_chunks(ss.chunks)
                    else:
                        query = standalone_question(llm, question, history)
                        docs = retrieve(query, top_k, use_mmr)
                context = format_context(docs)
            else:
                context = "FREE CHAT MODE" if has_docs else "NO DOCUMENTS UPLOADED"

            chain = PROMPT | llm | StrOutputParser()
            result = {"mode": "CHAT"}

            def on_mode(md: str) -> None:
                result["mode"] = md
                badge_ph.markdown(MODE_BADGES[md], unsafe_allow_html=True)

            answer = st.write_stream(
                stream_answer(chain, {"context": context, "history": lc_history, "question": question}, on_mode))
        except Exception as e:  # noqa: BLE001
            ss.messages.pop()
            st.error(f"Something went wrong: {str(e)[:300]}")
            return

        sources = []
        if use_docs and result["mode"] == "DOC":
            sources = [{"label": doc_label(d), "text": d.page_content[:500]} for d in docs]
            render_sources(sources)

    ss.messages.append({"role": "assistant", "content": answer, "mode": result["mode"], "sources": sources})


# ---------------------------------------------------
# Main tabs
# ---------------------------------------------------
tab_chat, tab_lib = st.tabs(["💬 Chat", "📚 Library & Stats"])

clicked = None
with tab_chat:
    if has_docs and mode.startswith("Smart"):
        cols = st.columns(len(QUICK_ACTIONS))
        for col, (label, prompt_text) in zip(cols, QUICK_ACTIONS.items()):
            if col.button(label, key=f"qa_{label}"):
                clicked = prompt_text
    elif not has_docs:
        st.info("👆 Upload files to chat with them — or just start chatting, I work as a regular ChatGPT too.")
    else:
        st.info("💬 Free chat mode is on — your documents are ignored until you switch back to Smart mode.")

    for m in ss.messages:
        render_message(m)

with tab_lib:
    if not has_docs:
        st.info("No files indexed yet.")
    else:
        total_chars = sum(len(c.page_content) for c in ss.chunks)
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Files", len(ss.file_names))
        c2.metric("Chunks", len(ss.chunks))
        c3.metric("Characters", f"{total_chars:,}")
        c4.metric("≈ Tokens", f"{total_chars // 4:,}")
        st.markdown("".join(f'<span class="filechip">📎 {n}</span>' for n in ss.file_names), unsafe_allow_html=True)
        st.write("")
        for f in uploaded or []:
            if f.name.rsplit(".", 1)[-1].lower() in IMAGE_EXT:
                st.image(f.getvalue(), caption=f.name, width=320)
        with st.expander("🧩 Preview of indexed chunks"):
            for c in ss.chunks[:8]:
                st.markdown(f'<div class="src-title">{doc_label(c)}</div>', unsafe_allow_html=True)
                st.text(c.page_content[:300])
                st.divider()

question = st.chat_input("Ask about your files — or anything else…")
with tab_chat:
    if clicked:
        handle_question(clicked, wide=True)
    elif question:
        handle_question(question)


# ---------------------------------------------------
# Sidebar chat-history list (rendered last so titles are up to date)
# ---------------------------------------------------
with hist_box:
    visible = [(cid, c) for cid, c in reversed(list(ss.chats.items())) if c["messages"]]
    if not visible:
        st.caption("No conversations yet — ask something!")
    for cid, c in visible:
        title = c["title"] if len(c["title"]) <= 30 else c["title"][:29] + "…"
        col_open, col_del = st.columns([5, 1])
        col_open.button(("🟢 " if cid == ss.active else "💬 ") + title, key=f"open_{cid}",
                        on_click=switch_chat, args=(cid,))
        col_del.button("🗑️", key=f"del_{cid}", on_click=delete_chat, args=(cid,))